import psycopg2
import os
import pandas as pd
from dotenv import load_dotenv
from openpyxl import Workbook
from openpyxl.utils.dataframe import dataframe_to_rows
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side

# Explicitly load the .env file from the correct directory
dotenv_path = os.path.join(os.path.expanduser('~'), 'grantseeker-azure-production', '.env')
if os.path.exists(dotenv_path):
    load_dotenv(dotenv_path)
else:
    load_dotenv()

print("=" * 80)
print("📊 EXPORTING OPPORTUNITIES TO EXCEL")
print("=" * 80)

# Connect to DB
try:
    conn = psycopg2.connect(
        host=os.getenv('DB_HOST'),
        user=os.getenv('DB_USER'),
        password=os.getenv('DB_PASSWORD'),
        database=os.getenv('DB_NAME'),
        port=os.getenv('DB_PORT'),
        sslmode='require'
    )
    print("✅ Database connected.")
except Exception as e:
    print(f"❌ Connection failed: {e}")
    exit(1)

# Query
query = """
    SELECT 
        f.name as "Foundation Name",
        f.website as "Website",
        f.is_grantmaking_charity as "Grantmaking?",
        f.funder_type_reason as "Classification Reason",
        fo.opportunity_title as "Opportunity Title",
        fo.funding_focus as "Funding Focus",
        fo.eligibility_inclusion as "Eligibility (Inclusion)",
        fo.eligibility_exclusion as "Eligibility (Exclusion)",
        fo.funding_amounts as "Funding Amounts",
        fo.deadlines as "Deadlines",
        fo.application_process as "Application Process",
        fo.application_form_type as "Form Type",
        fo.application_form_url as "Form URL",
        fo.application_questions_list as "Questions",
        fo.guidance_text as "Guidance",
        fo.guidance_url as "Guidance URL",
        fo.important_urls as "Important URLs",
        fo.created_at as "Date Captured"
    FROM funding_opportunities fo
    JOIN funders f ON fo.funder_id = f.id
    ORDER BY f.name, fo.opportunity_title
"""

print("⏳ Fetching data...")
df = pd.read_sql_query(query, conn)
conn.close()

print(f"✅ Fetched {len(df)} records.")

# Create Excel File with Formatting
output_file = "grant_opportunities_formatted.xlsx"
print(f"🎨 Creating formatted Excel file: {output_file}...")

wb = Workbook()
ws = wb.active
ws.title = "Opportunities"

# Convert DataFrame to rows
for r in dataframe_to_rows(df, index=False, header=True):
    ws.append(r)

# --- FORMATTING ---

# Define styles
header_font = Font(bold=True, color="FFFFFF")
header_fill = PatternFill(start_color="4F81BD", end_color="4F81BD", fill_type="solid")
border_style = Side(style='thin')
border = Border(left=border_style, right=border_style, top=border_style, bottom=border_style)
alignment_wrap = Alignment(wrap_text=True, vertical="top")

# Apply styles to header row
for cell in ws[1]:
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = Alignment(horizontal="center", vertical="center")

# Set column widths (approximate characters)
column_widths = {
    'A': 30,  # Foundation Name
    'B': 25,  # Website
    'C': 12,  # Grantmaking?
    'D': 30,  # Reason
    'E': 35,  # Opportunity Title
    'F': 40,  # Funding Focus
    'G': 40,  # Eligibility (Inclusion)
    'H': 40,  # Eligibility (Exclusion)
    'I': 20,  # Amounts
    'J': 20,  # Deadlines
    'K': 40,  # Process
    'L': 20,  # Form Type
    'M': 25,  # Form URL
    'N': 50,  # Questions
    'O': 50,  # Guidance
    'P': 25,  # Guidance URL
    'Q': 30,  # Important URLs
    'R': 20   # Date
}

for col_letter, width in column_widths.items():
    ws.column_dimensions[col_letter].width = width

# Apply text wrapping and alignment to all data cells
for row in ws.iter_rows(min_row=2):
    for cell in row:
        cell.alignment = alignment_wrap
        # Limit cell height indirectly? (Excel handles auto-height usually)
        
        # Truncate extremely long cells to prevent Excel crash (32k char limit)
        if isinstance(cell.value, str) and len(cell.value) > 32000:
            cell.value = cell.value[:32000] + "... (TRUNCATED)"

# Save
wb.save(output_file)

print("-" * 80)
print(f"✅ Export Complete!")
print(f"📁 File: {os.path.abspath(output_file)}")
print("-" * 80)